"""
Gera a planilha consolidada da Higienização de Rede 2.0 a partir das três bases:

  1. Faturamento 12 meses (CSV ; utf-8)          -> cpf_cnpj_prestador ; Soma de valorliberado
  2. Rede credenciada ERP FACIL (CSV ; cp1252)    -> cadastro completo do ERP
  3. Rede credenciada BI CODER (XLSX, aba Export) -> visão curada (categoria oficial, município IBGE)

Uso:
    python pipeline/gerar_consolidado.py --faturamento F.csv --facil E.csv --coder C.xlsx \
        --saida data/HIGIENIZACAO_REDE_CONSOLIDADA.xlsx [--data-ref 2026-09-18]
"""

import argparse
import re
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERDE = "007A33"
LIMA = "BBDD00"
CREME = "F0EDE4"

# Regras de escopo — nunca são candidatos a descredenciamento:
#   - enfermeiros e derivados (qualquer categoria com "ENFERM": ENFERMAGEM, ENFERMEIRO REGULADOR...)
#   - médico regulador
#   - rede própria Ipasgo (CLINICA IPASGO)
CATEGORIAS_EXCLUIDAS = {"MEDICO REGULADOR", "ENFERMEIRO REGULADOR", "CLINICA IPASGO"}
PADRAO_EXCLUIDO = re.compile(r"ENFERM", re.I)


def categoria_excluida(cat):
    cat = "" if cat is None or (isinstance(cat, float) and cat != cat) else str(cat)
    return cat in CATEGORIAS_EXCLUIDAS or bool(PADRAO_EXCLUIDO.search(cat))
NOMES_TECNICOS = re.compile(r"LIB\. SEM LOCAL|REEMBOLSO", re.I)

LEGENDA_BLOCOS = {
    "A": "Fatura nos últimos 12 meses — manter vínculo",
    "B": "Sem faturamento, credenciamento recente (< 12 meses) — aguardar maturação",
    "C": "Candidato a descredenciamento — rede suficiente na localidade (3+ efetivos sem ele)",
    "D": "Candidato — rede mínima (ELE + 1): descredenciar só com substituto confirmado",
    "E": "Candidato — ÚNICO na localidade/categoria: manter ou substituir antes (RN 567/585)",
    "F": "Candidato com EFEITO CUMULATIVO — vários da mesma localidade/categoria; escalonar lotes",
    "G": "Fatura mas não consta no BI CODER — divergência cadastral a corrigir",
    "H": "Sem faturamento e presente em só uma base cadastral (FACIL × CODER) — saneamento cadastral",
    "I": "Fora do escopo — enfermeiros e derivados, médico regulador, rede própria Ipasgo, registro técnico ou sem documento",
}


# ---------------------------------------------------------------- utilidades
def so_digitos(v):
    if pd.isna(v):
        return ""
    return re.sub(r"\D", "", str(v))


def valor_br(v):
    s = str(v).replace("R$", "").replace(".", "").replace(",", ".").strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def data_br(v, fmt):
    try:
        return datetime.strptime(str(v).strip(), fmt).date()
    except (ValueError, TypeError):
        return None


def classe_rede(n):
    if n <= 0:
        return "SEM REDE"
    if n == 1:
        return "ÚNICO NA LOCALIDADE"
    if n == 2:
        return "ELE + 1 (REDE MÍNIMA)"
    return "REDE SUFICIENTE"


# ---------------------------------------------------------------- leitura
def ler_coder(caminho):
    df = pd.read_excel(caminho, sheet_name="Export", dtype=str)
    df = df[df["PRESTADOR"].notna()].copy()
    ext = df["PRESTADOR"].str.extract(r"^(\d+)-(.*?)(?:\(.*\))?$")
    df["DOC"] = ext[0].fillna("")
    df["NOME_CODER"] = ext[1].str.strip()
    df["DOC"] = df["DOC"].apply(lambda d: d.zfill(11) if len(d) <= 11 else d.zfill(14))
    df = df.rename(
        columns={
            "MATRICULA": "MATRICULA_CODER",
            "DATA CREDENCIAMENTO": "DATA_CRED_CODER",
            "CATEGORIA": "CATEGORIA_CODER",
            "REGIÃO DE SAÚDE": "REGIAO_SAUDE",
            "IBGE_MUN": "IBGE_MUN",
            "MUNICÍPIO - IBGE": "MUNICIPIO",
            "Endereco_Completo": "ENDERECO",
            "TELEFONE": "TELEFONE",
        }
    )
    return df


def ler_facil(caminho):
    df = pd.read_csv(caminho, sep=";", dtype=str, encoding="cp1252")
    df["DOC"] = [
        (
            so_digitos(cnpj).zfill(14)
            if so_digitos(cnpj)
            else so_digitos(cpf).zfill(11) if so_digitos(cpf) else ""
        )
        for cpf, cnpj in zip(df["CPF"], df["CNPJ"])
    ]
    df["TECNICO"] = df["NOME"].fillna("").str.contains(NOMES_TECNICOS)
    agg = (
        df.groupby("DOC", dropna=False)
        .agg(
            NOME_FACIL=("NOME", "first"),
            TIPO_ERP_FACIL=("TIPO", "first"),
            CODIGO_FACIL=("CODIGO", lambda s: ", ".join(sorted(set(s.dropna())))),
            QTD_CADASTROS_FACIL=("CODIGO", "size"),
            DT_INCL_FACIL=("DT_INCL", "first"),
            DT_INI_CONTRATO=("DT_INI_CONTRATO", "first"),
            CNES=("CNES", "first"),
            EMAIL=("EMAIL", "first"),
            ENVIA_RPS_ANS=("ENVIA_RPS_ANS", "first"),
            URG_EMERGENCIA=("ID_URG_EMER", "first"),
            TECNICO=("TECNICO", "max"),
        )
        .reset_index()
    )
    return agg


def ler_faturamento(caminho, docs_conhecidos):
    df = pd.read_csv(caminho, sep=";", dtype=str, encoding="utf-8-sig")
    df.columns = ["DOC_BRUTO", "VALOR"]
    df["VALOR_FATURADO_12M"] = df["VALOR"].apply(valor_br)

    def normaliza(bruto):
        d = so_digitos(bruto)
        cpf, cnpj = d.zfill(11) if len(d) <= 11 else None, d.zfill(14)
        if cpf and cpf in docs_conhecidos:
            return cpf, "CPF"
        if cnpj in docs_conhecidos:
            return cnpj, "CNPJ"
        return (cpf, "CPF (sem cadastro)") if cpf else (cnpj, "CNPJ (sem cadastro)")

    res = df["DOC_BRUTO"].apply(normaliza)
    df["DOC"] = [r[0] for r in res]
    df["NORMALIZACAO_DOC"] = [r[1] for r in res]
    df["ZEROS_RESTAURADOS"] = df["DOC"].str.len() - df["DOC_BRUTO"].str.len()
    return (
        df.groupby("DOC")
        .agg(
            VALOR_FATURADO_12M=("VALOR_FATURADO_12M", "sum"),
            NORMALIZACAO_DOC=("NORMALIZACAO_DOC", "first"),
            ZEROS_RESTAURADOS=("ZEROS_RESTAURADOS", "max"),
        )
        .reset_index()
    )


# ---------------------------------------------------------------- consolidação
def consolidar(fat_csv, facil_csv, coder_xlsx, data_ref):
    coder = ler_coder(coder_xlsx)
    facil = ler_facil(facil_csv)
    conhecidos = set(coder["DOC"]) | set(facil["DOC"])
    fat = ler_faturamento(fat_csv, conhecidos)

    coder1 = coder.drop_duplicates("DOC")
    base = pd.DataFrame(
        {"DOC": sorted(set(coder1["DOC"]) | set(facil["DOC"]) | set(fat["DOC"]))}
    )
    base = (
        base.merge(coder1, on="DOC", how="left")
        .merge(facil, on="DOC", how="left")
        .merge(fat, on="DOC", how="left")
    )

    base["NO_CODER"] = base["MATRICULA_CODER"].notna()
    base["NO_FACIL"] = base["NOME_FACIL"].notna()
    base["VALOR_FATURADO_12M"] = base["VALOR_FATURADO_12M"].fillna(0.0)
    base["FATURA_12M"] = base["VALOR_FATURADO_12M"] > 0
    base["TECNICO"] = base["TECNICO"].fillna(False).astype(bool)
    base["TIPO_DOC"] = (
        base["DOC"].str.len().map({11: "CPF", 14: "CNPJ"}).fillna("SEM DOCUMENTO")
    )
    base.loc[base["DOC"] == "", "TIPO_DOC"] = "SEM DOCUMENTO"
    base["NOME"] = (
        base["NOME_CODER"]
        .fillna(base["NOME_FACIL"])
        .fillna("(não identificado — só no faturamento)")
    )

    def presenca(r):
        p = [
            n
            for n, ok in (
                ("CODER", r.NO_CODER),
                ("FACIL", r.NO_FACIL),
                ("FATURAMENTO", r.FATURA_12M),
            )
            if ok
        ]
        return " + ".join(p) if p else "NENHUMA"

    base["PRESENCA_BASES"] = base.apply(presenca, axis=1)

    # credenciamento recente: FACIL DT_INCL (dd/mm/aaaa) ou CODER (mm/aaaa)
    limite = date(data_ref.year - 1, data_ref.month, data_ref.day)

    def data_cred(r):
        d1 = data_br(r.DT_INCL_FACIL, "%d/%m/%Y")
        d2 = data_br(r.DATA_CRED_CODER, "%m/%Y")
        datas = [d for d in (d1, d2) if d]
        return max(datas) if datas else None

    base["DATA_CREDENCIAMENTO"] = base.apply(data_cred, axis=1)
    base["CREDENCIAMENTO_RECENTE"] = base["DATA_CREDENCIAMENTO"].apply(
        lambda d: bool(d and d >= limite)
    )
    base["ANOS_CREDENCIADO"] = base["DATA_CREDENCIAMENTO"].apply(
        lambda d: round((data_ref - d).days / 365.25, 1) if d else None
    )

    # disponibilidade por município (IBGE) × categoria oficial CODER
    loc = base[base["NO_CODER"]]
    cad = (
        loc.groupby(["IBGE_MUN", "CATEGORIA_CODER"]).size().rename("REDE_CADASTRAL_MUN")
    )
    efe = (
        loc[loc["FATURA_12M"]]
        .groupby(["IBGE_MUN", "CATEGORIA_CODER"])
        .size()
        .rename("REDE_EFETIVA_MUN")
    )
    reg = (
        loc[loc["FATURA_12M"]]
        .groupby(["REGIAO_SAUDE", "CATEGORIA_CODER"])
        .size()
        .rename("REDE_EFETIVA_REGIAO")
    )
    base = (
        base.join(cad, on=["IBGE_MUN", "CATEGORIA_CODER"])
        .join(efe, on=["IBGE_MUN", "CATEGORIA_CODER"])
        .join(reg, on=["REGIAO_SAUDE", "CATEGORIA_CODER"])
    )
    for c in ("REDE_CADASTRAL_MUN", "REDE_EFETIVA_MUN", "REDE_EFETIVA_REGIAO"):
        base[c] = base[c].fillna(0).astype(int)
    base["CLASSE_REDE_CADASTRAL"] = base["REDE_CADASTRAL_MUN"].apply(classe_rede)
    base["CLASSE_REDE_EFETIVA"] = base["REDE_EFETIVA_MUN"].apply(classe_rede)

    # candidatos (sem faturamento, não recentes, no CODER e no FACIL, fora de exclusões)
    excluido = (
        base["CATEGORIA_CODER"].apply(categoria_excluida)
        | base["TECNICO"]
        | (base["TIPO_DOC"] == "SEM DOCUMENTO")
    )
    base["CANDIDATO"] = (
        (~base["FATURA_12M"])
        & (~base["CREDENCIAMENTO_RECENTE"])
        & base["NO_CODER"]
        & base["NO_FACIL"]
        & ~excluido
    )

    # efeito cumulativo: remover todos os candidatos do grupo rebaixa a classe cadastral
    cand_grp = (
        base[base["CANDIDATO"]]
        .groupby(["IBGE_MUN", "CATEGORIA_CODER"])
        .size()
        .rename("CANDIDATOS_NO_GRUPO")
    )
    base = base.join(cand_grp, on=["IBGE_MUN", "CATEGORIA_CODER"])
    base["CANDIDATOS_NO_GRUPO"] = base["CANDIDATOS_NO_GRUPO"].fillna(0).astype(int)
    base["REDE_APOS_LOTE"] = base["REDE_CADASTRAL_MUN"] - base["CANDIDATOS_NO_GRUPO"]
    base["EFEITO_CUMULATIVO"] = (
        base["CANDIDATO"]
        & (base["CANDIDATOS_NO_GRUPO"] >= 2)
        & (base["REDE_APOS_LOTE"].apply(classe_rede) != base["CLASSE_REDE_CADASTRAL"])
        & (base["REDE_APOS_LOTE"] < 3)
    )

    # bloco de decisão
    def bloco(r):
        if (
            categoria_excluida(r.CATEGORIA_CODER)
            or r.TECNICO
            or r.TIPO_DOC == "SEM DOCUMENTO"
        ):
            return "I"
        if r.FATURA_12M:
            return "A" if r.NO_CODER else "G"
        if not (r.NO_CODER and r.NO_FACIL):
            return "H"
        if r.CREDENCIAMENTO_RECENTE:
            return "B"
        if r.EFEITO_CUMULATIVO:
            return "F"
        sem_ele = (
            r.REDE_EFETIVA_MUN
        )  # ele não fatura, então a rede efetiva já está "sem ele"
        if sem_ele >= 3:
            return "C"
        if sem_ele >= 1:
            return "D"
        return "E"

    base["BLOCO"] = base.apply(bloco, axis=1)
    base["DESCRICAO_BLOCO"] = base["BLOCO"].map(LEGENDA_BLOCOS)

    acao = {
        "A": "MANTER",
        "B": "MANTER — REAVALIAR NO PRÓXIMO CICLO",
        "C": "DESCREDENCIAR",
        "D": "DESCREDENCIAR COM SUBSTITUTO CONFIRMADO",
        "E": "MANTER / BUSCAR SUBSTITUTO ANTES",
        "F": "DESCREDENCIAR EM LOTES ESCALONADOS",
        "G": "REGULARIZAR CADASTRO NO BI CODER",
        "H": "SANEAR CADASTRO (FACIL × CODER)",
        "I": "FORA DO ESCOPO",
    }
    base["ACAO_RECOMENDADA"] = base["BLOCO"].map(acao)

    hospitalar = (
        base["CATEGORIA_CODER"].fillna("").str.contains("HOSPITAL")
        | (base["TIPO_ERP_FACIL"] == "HOSPITAL") & base["CATEGORIA_CODER"].isna()
    )
    base["NORMA_APLICAVEL"] = ""
    m = base["BLOCO"].isin(list("CDEF"))
    base.loc[m & hospitalar, "NORMA_APLICAVEL"] = (
        "RN 585/2023 (rede hospitalar) + Lei 13.003/2014"
    )
    base.loc[m & ~hospitalar, "NORMA_APLICAVEL"] = (
        "RN 567/2022 (não hospitalar) + Lei 13.003/2014"
    )

    # score 0-100 (prioridade de descredenciamento) com fatores explicados
    def score(r):
        if not r.CANDIDATO:
            return 0, "Não candidato"
        pts, fat = 40, ["+40 sem faturamento em 12 meses"]
        if r.ANOS_CREDENCIADO and r.ANOS_CREDENCIADO >= 5:
            pts += 15
            fat.append("+15 credenciado há 5+ anos")
        elif r.ANOS_CREDENCIADO and r.ANOS_CREDENCIADO >= 2:
            pts += 8
            fat.append("+8 credenciado há 2–5 anos")
        if r.REDE_EFETIVA_MUN >= 3:
            pts += 30
            fat.append("+30 rede efetiva suficiente no município")
        elif r.REDE_EFETIVA_MUN >= 1:
            pts += 12
            fat.append("+12 há rede efetiva mínima no município")
        if r.REDE_EFETIVA_REGIAO >= 3:
            pts += 10
            fat.append("+10 rede efetiva suficiente na região de saúde")
        if pd.isna(r.EMAIL) and pd.isna(r.TELEFONE):
            pts += 5
            fat.append("+5 sem contato cadastrado")
        if r.EFEITO_CUMULATIVO:
            pts -= 15
            fat.append("−15 efeito cumulativo no grupo")
        return max(0, min(100, pts)), "; ".join(fat)

    sc = base.apply(score, axis=1)
    base["SCORE_PRIORIDADE"] = [s[0] for s in sc]
    base["FATORES_SCORE"] = [s[1] for s in sc]

    def obs(r):
        o = []
        if r.QTD_CADASTROS_FACIL and r.QTD_CADASTROS_FACIL > 1:
            o.append(f"{int(r.QTD_CADASTROS_FACIL)} cadastros no FACIL")
        if (
            isinstance(r.NORMALIZACAO_DOC, str)
            and r.ZEROS_RESTAURADOS
            and r.ZEROS_RESTAURADOS > 0
        ):
            o.append(
                f"doc do faturamento normalizado ({int(r.ZEROS_RESTAURADOS)} zero(s) restaurado(s))"
            )
        if isinstance(r.NORMALIZACAO_DOC, str) and "sem cadastro" in r.NORMALIZACAO_DOC:
            o.append("fatura mas não está em nenhuma base cadastral")
        if r.NO_CODER and not r.NO_FACIL:
            o.append("consta no CODER, ausente no FACIL")
        if r.NO_FACIL and not r.NO_CODER:
            o.append("consta no FACIL, ausente no CODER")
        if r.BLOCO == "E":
            o.append(
                "ausência de faturamento sozinha não autoriza descredenciamento (exige substituto)"
            )
        return "; ".join(o)

    base["OBSERVACOES"] = base.apply(obs, axis=1)

    base["DATA_CREDENCIAMENTO"] = base["DATA_CREDENCIAMENTO"].apply(
        lambda d: d.strftime("%d/%m/%Y") if d else ""
    )
    base["STATUS_ACOMPANHAMENTO"] = "PENDENTE"
    base["RESPONSAVEL"] = ""

    colunas = [
        "DOC",
        "TIPO_DOC",
        "NOME",
        "CATEGORIA_CODER",
        "TIPO_ERP_FACIL",
        "UF",
        "REGIAO_SAUDE",
        "IBGE_MUN",
        "MUNICIPIO",
        "BAIRRO",
        "ENDERECO",
        "TELEFONE",
        "EMAIL",
        "CNES",
        "MATRICULA_CODER",
        "CODIGO_FACIL",
        "QTD_CADASTROS_FACIL",
        "DATA_CREDENCIAMENTO",
        "ANOS_CREDENCIADO",
        "CREDENCIAMENTO_RECENTE",
        "PRESENCA_BASES",
        "VALOR_FATURADO_12M",
        "FATURA_12M",
        "REDE_CADASTRAL_MUN",
        "REDE_EFETIVA_MUN",
        "REDE_EFETIVA_REGIAO",
        "CLASSE_REDE_CADASTRAL",
        "CLASSE_REDE_EFETIVA",
        "CANDIDATOS_NO_GRUPO",
        "EFEITO_CUMULATIVO",
        "BLOCO",
        "DESCRICAO_BLOCO",
        "ACAO_RECOMENDADA",
        "NORMA_APLICAVEL",
        "SCORE_PRIORIDADE",
        "FATORES_SCORE",
        "OBSERVACOES",
        "ENVIA_RPS_ANS",
        "STATUS_ACOMPANHAMENTO",
        "RESPONSAVEL",
    ]
    out = base[colunas].copy()
    for c in ("CREDENCIAMENTO_RECENTE", "FATURA_12M", "EFEITO_CUMULATIVO"):
        out[c] = out[c].map({True: "SIM", False: "NÃO"})
    out["QTD_CADASTROS_FACIL"] = out["QTD_CADASTROS_FACIL"].fillna(0).astype(int)
    out = out.fillna("")
    return out.sort_values(
        ["BLOCO", "SCORE_PRIORIDADE"], ascending=[True, False]
    ).reset_index(drop=True)


# ---------------------------------------------------------------- saída
def escrever(df, saida, data_ref):
    resumo = (
        df.groupby(["BLOCO", "DESCRICAO_BLOCO", "ACAO_RECOMENDADA"])
        .agg(
            PRESTADORES=("DOC", "size"),
            VALOR_FATURADO_12M=("VALOR_FATURADO_12M", "sum"),
        )
        .reset_index()
    )
    disp = (
        df[df["CATEGORIA_CODER"] != ""]
        .groupby(["UF", "REGIAO_SAUDE", "MUNICIPIO", "CATEGORIA_CODER"])
        .agg(
            REDE_CADASTRAL=("DOC", "size"),
            REDE_EFETIVA=("FATURA_12M", lambda s: (s == "SIM").sum()),
            CANDIDATOS=("BLOCO", lambda s: s.isin(list("CDEF")).sum()),
        )
        .reset_index()
    )
    legenda = pd.DataFrame(
        [{"BLOCO": k, "DESCRICAO": v} for k, v in LEGENDA_BLOCOS.items()]
    )
    info = pd.DataFrame(
        {
            "ITEM": ["Data de referência", "Gerado em", "Total de prestadores"],
            "VALOR": [
                data_ref.strftime("%d/%m/%Y"),
                datetime.now().strftime("%d/%m/%Y %H:%M"),
                len(df),
            ],
        }
    )

    with pd.ExcelWriter(saida, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="CONSOLIDADO", index=False)
        resumo.to_excel(w, sheet_name="RESUMO", index=False)
        disp.to_excel(w, sheet_name="DISPONIBILIDADE", index=False)
        legenda.to_excel(w, sheet_name="LEGENDA", index=False)
        info.to_excel(w, sheet_name="INFO", index=False)

    wb = load_workbook(saida)
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=VERDE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(
            ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), 1
        ):
            largura = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(
                max(10, largura + 2), 50
            )
    ws = wb["CONSOLIDADO"]
    col_valor = [c.column for c in ws[1] if c.value == "VALOR_FATURADO_12M"][0]
    for row in ws.iter_rows(min_row=2, min_col=col_valor, max_col=col_valor):
        row[0].number_format = '"R$" #,##0.00'
    wb.save(saida)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--faturamento", required=True)
    ap.add_argument("--facil", required=True)
    ap.add_argument("--coder", required=True)
    ap.add_argument("--saida", required=True)
    ap.add_argument("--data-ref", default="2026-09-18")
    a = ap.parse_args()
    ref = datetime.strptime(a.data_ref, "%Y-%m-%d").date()
    df = consolidar(a.faturamento, a.facil, a.coder, ref)
    escrever(df, a.saida, ref)
    print(f"{len(df)} linhas × {len(df.columns)} colunas -> {a.saida}")
    print(df["BLOCO"].value_counts().sort_index().to_string())
